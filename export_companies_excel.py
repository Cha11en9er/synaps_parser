"""
Общий companies.json → Excel.

Используется из ritual_comp / developer_comp / genpodryad_comp.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

COLUMNS: list[tuple[str, str, int]] = [
    ("inn", "ИНН", 14),
    ("name", "Название", 36),
    ("status", "Статус", 10),
    ("profile_url", "Ссылка Synaps", 28),
    ("registration_date", "Дата регистрации", 14),
    ("authorized_capital", "Уставный капитал", 16),
    ("phones", "Телефоны", 22),
    ("emails", "Почта", 28),
    ("okved", "ОКВЭД", 32),
    ("legal_address", "Адрес", 32),
    ("revenue_2023", "Выручка 2023", 14),
    ("revenue_2024", "Выручка 2024", 14),
    ("revenue_2025", "Выручка 2025", 14),
    ("revenue_trend", "Динамика выручки", 16),
    ("reliability", "Надёжность", 12),
    ("enforcement_debt", "Исп. производство", 28),
    ("general_director", "Руководитель", 28),
    ("headcount", "Численность", 12),
    ("bank_account_state", "Банк. счета", 12),
    ("error", "Ошибка", 20),
]

HEADER_FONT = Font(bold=True)
INDEX_FONT = Font(bold=True)
NO_BORDER = Border(
    left=Side(style=None),
    right=Side(style=None),
    top=Side(style=None),
    bottom=Side(style=None),
)
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
HEADER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")
PT_PER_LINE = 15.0
MAX_ROW_HEIGHT = 409.0


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _line_count(text: str) -> int:
    if not text:
        return 1
    lines = 0
    for part in text.split("\n"):
        lines += max(1, (len(part) // 40) + (1 if part else 0)) if part else 1
    return max(1, lines)


def load_names(*paths: Path) -> dict[str, str]:
    out: dict[str, str] = {}
    for path in paths:
        if not path or not path.is_file():
            continue
        raw = json.loads(path.read_text(encoding="utf-8"))
        pools: list[Any] = []
        if isinstance(raw, dict):
            for key in ("companies", "items", "names"):
                if isinstance(raw.get(key), list):
                    pools.append(raw[key])
        elif isinstance(raw, list):
            pools.append(raw)
        for pool in pools:
            for item in pool:
                if isinstance(item, str):
                    continue
                if not isinstance(item, dict):
                    continue
                inn = re.sub(r"\D", "", str(item.get("inn") or ""))
                name = _as_text(
                    item.get("name")
                    or item.get("short_name")
                    or item.get("full_name")
                    or item.get("query_name")
                    or item.get("title")
                    or ""
                )
                if inn and name and inn not in out:
                    out[inn] = name
    return out


def flatten_item(item: dict[str, Any], names: dict[str, str]) -> dict[str, str]:
    inn = re.sub(r"\D", "", str(item.get("inn") or ""))
    data = item.get("data") if isinstance(item.get("data"), dict) else {}
    name = _as_text(item.get("name") or item.get("query_name") or "") or names.get(inn, "")
    row: dict[str, str] = {
        "inn": inn,
        "name": name,
        "status": _as_text(item.get("status")),
        "profile_url": _as_text(item.get("profile_url")),
        "error": _as_text(item.get("error")),
    }
    for key, _, _ in COLUMNS:
        if key in row:
            continue
        row[key] = _as_text(data.get(key))
    return row


def unique_output_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem, suffix, parent = path.stem, path.suffix or ".xlsx", path.parent
    n = 2
    while True:
        candidate = parent / f"{stem}_{n}{suffix}"
        if not candidate.exists():
            return candidate
        n += 1


def export_excel(
    input_path: Path,
    output_path: Path,
    *,
    sheet_title: str,
    name_sources: list[Path] | None = None,
    overwrite: bool = False,
) -> Path:
    payload = json.loads(input_path.read_text(encoding="utf-8"))
    items = payload.get("items") if isinstance(payload, dict) else None
    if not isinstance(items, list):
        raise SystemExit(f"В {input_path} нет списка items")

    if not overwrite:
        output_path = unique_output_path(output_path)

    names = load_names(*(name_sources or []))
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Лист1"

    headers = ["№"] + [title for _, title, _ in COLUMNS]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(1, col_idx)
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
        cell.border = NO_BORDER

    for i, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            continue
        flat = flatten_item(item, names)
        values = [i] + [flat.get(key, "") for key, _, _ in COLUMNS]
        ws.append(values)
        excel_row = ws.max_row

        max_lines = 1
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(excel_row, col_idx)
            cell.alignment = WRAP
            cell.border = NO_BORDER
            if col_idx == 1:
                cell.font = INDEX_FONT
                cell.alignment = Alignment(vertical="top", horizontal="center", wrap_text=True)
            max_lines = max(max_lines, _line_count(_as_text(value)))

        ws.row_dimensions[excel_row].height = min(MAX_ROW_HEIGHT, PT_PER_LINE * max_lines)

    ws.column_dimensions["A"].width = 6
    for i, (_, _, width) in enumerate(COLUMNS, start=2):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
